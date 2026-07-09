"""Report generator — Phase 7.

Produces PDF (ReportLab) and Excel (openpyxl) exports of a Load record.
Files are stored under EXPORT_DIR/{load_id}/ and served as file downloads.
"""
from __future__ import annotations

import io
import os
import uuid
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session, joinedload

from app.core.config import settings
from app.models.load import Load
from app.models.pallet import Pallet

# ── Brand colours ─────────────────────────────────────────────────────────────
_PURPLE = (123, 31, 162)   # #7B1FA2
_BG     = (250, 245, 252)  # #FAF5FC
_PASS   = (46, 125, 50)    # green
_HOLD   = (237, 108, 2)    # orange
_REJECT = (211, 47, 47)    # red
_GREY   = (117, 117, 117)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _export_dir(load_id: int) -> Path:
    path = Path(settings.EXPORT_DIR) / str(load_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _fmt(v: Optional[Decimal | float | int], unit: str = "") -> str:
    if v is None:
        return "—"
    return f"{v}{(' ' + unit) if unit else ''}"


def _status_color_hex(status: str) -> str:
    m = {"PASS": "#2e7d32", "HOLD": "#ed6c02", "REJECT": "#d32f2f"}
    return m.get(status, "#757575")


def _date_str(d: Optional[date | str]) -> str:
    if d is None:
        return "—"
    if isinstance(d, str):
        return d
    return d.strftime("%d %B %Y")


# ── PDF ───────────────────────────────────────────────────────────────────────

def generate_pdf(db: Session, load_id: int) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
    )

    load = _load_with_relations(db, load_id)
    out_path = _export_dir(load_id) / f"load_{load_id}_{uuid.uuid4().hex[:8]}.pdf"

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    brand = colors.Color(*[c / 255 for c in _PURPLE])
    pass_c = colors.Color(*[c / 255 for c in _PASS])
    hold_c = colors.Color(*[c / 255 for c in _HOLD])
    rej_c  = colors.Color(*[c / 255 for c in _REJECT])

    def status_color(s: str):
        return {
            "PASS": pass_c,
            "HOLD": hold_c,
            "REJECT": rej_c,
        }.get(s, colors.grey)

    title_style = ParagraphStyle("title", parent=styles["Heading1"],
                                  textColor=brand, fontSize=18, spaceAfter=4)
    h2_style    = ParagraphStyle("h2", parent=styles["Heading2"],
                                  textColor=brand, fontSize=11, spaceBefore=12, spaceAfter=4)
    normal      = styles["Normal"]
    small       = ParagraphStyle("small", parent=normal, fontSize=8, textColor=colors.grey)

    story = []

    # ── Title block ───────────────────────────────────────────────────────────
    story.append(Paragraph("MAGOPCO", title_style))
    story.append(Paragraph("Quality Inspection Report", styles["Heading2"]))
    story.append(HRFlowable(width="100%", color=brand, thickness=2, spaceAfter=6))

    # ── Load header ───────────────────────────────────────────────────────────
    header_data = [
        ["Load Reference",  load.load_reference or "—",
         "Inspection Date",  _date_str(load.inspection_date)],
        ["Container",       load.container_number or "—",
         "Inspection Place", load.inspection_place or "—"],
        ["Total Pallets",   str(load.total_pallets or "—"),
         "Total Cases",      str(load.total_cases or "—")],
    ]
    t = Table(header_data, colWidths=[3.5*cm, 5.5*cm, 3.5*cm, 5.5*cm])
    t.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (-1, -1), "Helvetica"),
        ("FONTNAME",  (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME",  (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.Color(0.98, 0.97, 0.99)]),
        ("GRID",      (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING",    (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    # ── Score summary ─────────────────────────────────────────────────────────
    story.append(Paragraph("Overall Result", h2_style))
    q  = f"Q{int(load.quality_score)}"  if load.quality_score  else "—"
    cs = f"CS-{load.condition_score}"   if load.condition_score else "—"
    score_data = [
        ["Final Status",    load.final_status],
        ["Quality Score",   q],
        ["Condition Score", cs],
        ["Main Issue",      load.main_issue or "None"],
    ]
    t = Table(score_data, colWidths=[4*cm, 14*cm])
    t.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (-1, -1), "Helvetica"),
        ("FONTNAME",  (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (1, 0), (1, 0), status_color(load.final_status)),
        ("FONTNAME",  (1, 0), (1, 0), "Helvetica-Bold"),
        ("GRID",      (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING",    (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.3*cm))

    # ── Summary measurements ──────────────────────────────────────────────────
    if load.summary_measurements:
        story.append(Paragraph("Average Measurements (Load Level)", h2_style))
        headers = ["Parameter", "Average", "Min", "Max", "Std Min", "Std Max", "Unit", "Result"]
        rows = [headers]
        for s in load.summary_measurements:
            rows.append([
                s.parameter_name,
                _fmt(s.average_value),
                _fmt(s.min_value),
                _fmt(s.max_value),
                _fmt(s.standard_min),
                _fmt(s.standard_max),
                s.unit or "—",
                s.status,
            ])
        t = Table(rows, repeatRows=1)
        style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), brand),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("GRID",       (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.98, 0.97, 0.99)]),
            ("ALIGN",      (1, 0), (-1, -1), "CENTER"),
            ("TOPPADDING",    (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i, row in enumerate(rows[1:], 1):
            if row[-1] == "FAIL":
                style_cmds.append(("TEXTCOLOR", (-1, i), (-1, i), rej_c))
                style_cmds.append(("FONTNAME",  (-1, i), (-1, i), "Helvetica-Bold"))
            elif row[-1] == "PASS":
                style_cmds.append(("TEXTCOLOR", (-1, i), (-1, i), pass_c))
        t.setStyle(TableStyle(style_cmds))
        story.append(t)

    # ── Per-pallet breakdown ──────────────────────────────────────────────────
    if load.pallets:
        story.append(Paragraph("Pallet Breakdown", h2_style))
        for pallet in load.pallets:
            story.append(Paragraph(
                f"<b>{pallet.pallet_number}</b> &nbsp;·&nbsp; "
                f"Status: <font color='{_status_color_hex(pallet.status)}'><b>{pallet.status}</b></font>"
                f"{(' &nbsp;· Q' + str(pallet.q_score)) if pallet.q_score else ''}"
                f"{(' &nbsp;· CS-' + pallet.cs_score) if pallet.cs_score else ''}",
                styles["Normal"],
            ))
            if pallet.measurements:
                p_headers = ["Parameter", "Value", "Std Min", "Std Max", "Dev.", "Severity", "Result"]
                p_rows = [p_headers]
                for m in pallet.measurements:
                    if m.value_numeric is None:
                        continue
                    dev = ""
                    if m.deviation is not None and m.deviation != 0:
                        dev = f"{'+' if m.deviation > 0 else ''}{m.deviation:.2f}"
                    p_rows.append([
                        m.parameter_name,
                        _fmt(m.value_numeric, m.unit or ""),
                        _fmt(m.standard_min),
                        _fmt(m.standard_max),
                        dev,
                        m.severity or "—",
                        m.status,
                    ])
                if len(p_rows) > 1:
                    pt = Table(p_rows, repeatRows=1)
                    ps = [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.Color(*[c / 255 for c in _BG])),
                        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE",   (0, 0), (-1, -1), 7.5),
                        ("GRID",       (0, 0), (-1, -1), 0.5, colors.lightgrey),
                        ("ALIGN",      (1, 0), (-1, -1), "CENTER"),
                        ("TOPPADDING",    (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                    for i, row in enumerate(p_rows[1:], 1):
                        if row[-1] == "FAIL":
                            ps.append(("BACKGROUND", (0, i), (-1, i), colors.Color(1, 0.96, 0.96)))
                            ps.append(("TEXTCOLOR",  (-1, i), (-1, i), rej_c))
                            ps.append(("FONTNAME",   (-1, i), (-1, i), "Helvetica-Bold"))
                        elif row[-1] == "PASS":
                            ps.append(("TEXTCOLOR", (-1, i), (-1, i), pass_c))
                    pt.setStyle(TableStyle(ps))
                    story.append(pt)
            story.append(Spacer(1, 0.3*cm))

    # ── Footer ────────────────────────────────────────────────────────────────
    story.append(HRFlowable(width="100%", color=colors.lightgrey, spaceAfter=4))
    story.append(Paragraph(
        f"Generated by Quality Intelligence Platform · {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC",
        small,
    ))

    doc.build(story)
    return out_path


# ── Excel ─────────────────────────────────────────────────────────────────────

def generate_xlsx(db: Session, load_id: int) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers,
    )
    from openpyxl.utils import get_column_letter

    load = _load_with_relations(db, load_id)
    out_path = _export_dir(load_id) / f"load_{load_id}_{uuid.uuid4().hex[:8]}.xlsx"

    wb = Workbook()

    # ── colour helpers ────────────────────────────────────────────────────────
    purple_fill = PatternFill("solid", fgColor="7B1FA2")
    bg_fill     = PatternFill("solid", fgColor="FAF5FC")
    pass_fill   = PatternFill("solid", fgColor="E8F5E9")
    fail_fill   = PatternFill("solid", fgColor="FFEBEE")
    hold_fill   = PatternFill("solid", fgColor="FFF8E1")
    header_font = Font(bold=True, color="FFFFFF")
    label_font  = Font(bold=True)
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    center      = Alignment(horizontal="center", vertical="center")

    status_fill = {"PASS": pass_fill, "FAIL": fail_fill, "HOLD": hold_fill}

    def write_header(ws, row, cols):
        for c, val in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=val)
            cell.font = header_font
            cell.fill = purple_fill
            cell.alignment = center
            cell.border = border

    def auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # ── Sheet 1: Load Summary ─────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Load Summary"

    rows = [
        ("Load Reference",  load.load_reference or "—"),
        ("Inspection Date",  _date_str(load.inspection_date)),
        ("Inspection Place", load.inspection_place or "—"),
        ("Container",        load.container_number or "—"),
        ("Vessel",           load.vessel_name or "—"),
        ("Total Pallets",    load.total_pallets),
        ("Total Cases",      load.total_cases),
        ("Total Weight",     float(load.total_weight) if load.total_weight else None),
        ("", ""),
        ("Final Status",    load.final_status),
        ("Quality Score",   f"Q{int(load.quality_score)}" if load.quality_score else "—"),
        ("Condition Score", f"CS-{load.condition_score}" if load.condition_score else "—"),
        ("Main Issue",      load.main_issue or "None"),
    ]
    for r, (label, val) in enumerate(rows, 1):
        ws1.cell(r, 1, label).font = label_font
        cell = ws1.cell(r, 2, val)
        if label == "Final Status":
            f = status_fill.get(str(val), None)
            if f:
                cell.fill = f
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 30

    # ── Sheet 2: Summary Measurements ────────────────────────────────────────
    ws2 = wb.create_sheet("Summary Measurements")
    cols2 = ["Parameter", "Average", "Min", "Max", "Std Min", "Std Max", "Unit", "Status", "Severity"]
    write_header(ws2, 1, cols2)
    for r, s in enumerate(load.summary_measurements, 2):
        ws2.cell(r, 1, s.parameter_name)
        ws2.cell(r, 2, float(s.average_value) if s.average_value is not None else None)
        ws2.cell(r, 3, float(s.min_value) if s.min_value is not None else None)
        ws2.cell(r, 4, float(s.max_value) if s.max_value is not None else None)
        ws2.cell(r, 5, float(s.standard_min) if s.standard_min is not None else None)
        ws2.cell(r, 6, float(s.standard_max) if s.standard_max is not None else None)
        ws2.cell(r, 7, s.unit or "—")
        cell = ws2.cell(r, 8, s.status)
        f = status_fill.get(s.status, None)
        if f:
            cell.fill = f
        ws2.cell(r, 9, s.severity or "—")
        for c in range(1, 10):
            ws2.cell(r, c).border = border
    auto_width(ws2)

    # ── Sheet 3: Pallet Details ───────────────────────────────────────────────
    ws3 = wb.create_sheet("Pallet Details")
    cols3 = [
        "Pallet #", "Grower Code", "GGN", "Packing Date", "Cases", "Weight",
        "Q Score", "CS Score", "Status",
        "Parameter", "Value", "Unit", "Std Min", "Std Max", "Deviation", "Severity", "Result",
    ]
    write_header(ws3, 1, cols3)
    r = 2
    for pallet in load.pallets:
        for m in pallet.measurements:
            ws3.cell(r, 1, pallet.pallet_number)
            ws3.cell(r, 2, pallet.grower_code or "—")
            ws3.cell(r, 3, pallet.ggn or "—")
            ws3.cell(r, 4, str(pallet.packing_date) if pallet.packing_date else "—")
            ws3.cell(r, 5, pallet.cases_count)
            ws3.cell(r, 6, float(pallet.weight) if pallet.weight else None)
            ws3.cell(r, 7, f"Q{pallet.q_score}" if pallet.q_score else "—")
            ws3.cell(r, 8, f"CS-{pallet.cs_score}" if pallet.cs_score else "—")
            cell_status = ws3.cell(r, 9, pallet.status)
            f = status_fill.get(pallet.status, None)
            if f:
                cell_status.fill = f
            ws3.cell(r, 10, m.parameter_name)
            ws3.cell(r, 11, float(m.value_numeric) if m.value_numeric is not None else m.value_text or "—")
            ws3.cell(r, 12, m.unit or "—")
            ws3.cell(r, 13, float(m.standard_min) if m.standard_min is not None else None)
            ws3.cell(r, 14, float(m.standard_max) if m.standard_max is not None else None)
            ws3.cell(r, 15, float(m.deviation) if m.deviation is not None else None)
            ws3.cell(r, 16, m.severity or "—")
            cell_res = ws3.cell(r, 17, m.status)
            rf = status_fill.get(m.status, None)
            if rf:
                cell_res.fill = rf
            for c in range(1, 18):
                ws3.cell(r, c).border = border
            r += 1
    auto_width(ws3)

    wb.save(str(out_path))
    return out_path


# ── Analytics PDF ─────────────────────────────────────────────────────────────

def generate_analytics_pdf(
    db: Session,
    months: int = 12,
    client_id: Optional[int] = None,
    period: str = "monthly",
) -> Path:
    from datetime import timedelta
    from sqlalchemy import case, func, String
    from app.models.pallet import Pallet
    from app.models.pallet_measurement import PalletMeasurement
    from app.models.client import Client
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

    # ── Fetch data using same logic as analytics endpoints ─────────────────────
    cutoff = date.today() - timedelta(days=months * 30)

    load_q = db.query(Load).filter(
        Load.inspection_date >= cutoff,
        Load.inspection_date.isnot(None),
    )
    client_name = "All Clients"
    if client_id:
        load_q = load_q.filter(Load.client_id == client_id)
        cl = db.get(Client, client_id)
        if cl:
            client_name = cl.name
    load_ids = [lo.id for lo in load_q.with_entities(Load.id).all()]

    # Overview KPIs
    passed_total = failed_total = held_total = total_pallets = 0
    pass_rate = avg_qs = worst_param = worst_rate = None
    if load_ids:
        from sqlalchemy import case as sa_case
        pa = (
            db.query(
                func.count(Pallet.id).label("total"),
                func.sum(sa_case((Pallet.status == "PASS", 1), else_=0)).label("passed"),
                func.sum(sa_case((Pallet.status.in_(["FAIL", "REJECT"]), 1), else_=0)).label("failed"),
                func.sum(sa_case((Pallet.status == "HOLD", 1), else_=0)).label("held"),
            )
            .filter(Pallet.load_id.in_(load_ids))
            .one()
        )
        total_pallets = pa.total or 0
        passed_total  = pa.passed or 0
        failed_total  = pa.failed or 0
        held_total    = pa.held   or 0
        analysed = passed_total + failed_total + held_total
        pass_rate = round(passed_total / analysed * 100, 1) if analysed else None

        avg_qs = (
            db.query(func.avg(Load.quality_score))
            .filter(Load.id.in_(load_ids), Load.quality_score.isnot(None))
            .scalar()
        )

        param_agg = (
            db.query(
                PalletMeasurement.parameter_name,
                func.count(PalletMeasurement.id).label("total"),
                func.sum(sa_case((PalletMeasurement.status == "PASS", 1), else_=0)).label("passed"),
            )
            .join(Pallet, PalletMeasurement.pallet_id == Pallet.id)
            .filter(Pallet.load_id.in_(load_ids), PalletMeasurement.status.in_(["PASS", "FAIL"]))
            .group_by(PalletMeasurement.parameter_name)
            .having(func.count(PalletMeasurement.id) >= 3)
            .all()
        )
        for row in param_agg:
            cr = round(row.passed / row.total * 100, 1) if row.total else None
            if cr is not None and (worst_rate is None or cr < worst_rate):
                worst_rate = cr
                worst_param = row.parameter_name

    # Quality trends
    trunc_expr = func.date_trunc("month" if period == "monthly" else "week", Load.inspection_date)
    pallet_trend = (
        db.query(
            trunc_expr.cast(String).label("period"),
            func.count(Pallet.id).label("total"),
            func.sum(case((Pallet.status == "PASS", 1), else_=0)).label("passed"),
            func.sum(case((Pallet.status.in_(["FAIL", "REJECT"]), 1), else_=0)).label("failed"),
            func.sum(case((Pallet.status == "HOLD", 1), else_=0)).label("held"),
        )
        .join(Load, Pallet.load_id == Load.id)
        .filter(Load.inspection_date >= cutoff, Load.inspection_date.isnot(None),
                Pallet.status.isnot(None))
        .group_by("period").order_by("period")
    )
    if client_id:
        pallet_trend = pallet_trend.filter(Load.client_id == client_id)
    trend_rows = pallet_trend.all()

    # Parameter compliance
    compliance_rows = []
    if load_ids:
        compliance_rows = (
            db.query(
                PalletMeasurement.parameter_name,
                PalletMeasurement.unit,
                func.count(PalletMeasurement.id).label("total"),
                func.sum(case((PalletMeasurement.status == "PASS", 1), else_=0)).label("passed"),
                func.avg(PalletMeasurement.value_numeric).label("avg_value"),
                func.max(PalletMeasurement.standard_max).label("std_max"),
            )
            .join(Pallet, PalletMeasurement.pallet_id == Pallet.id)
            .filter(
                Pallet.load_id.in_(load_ids),
                PalletMeasurement.status.in_(["PASS", "FAIL"]),
                PalletMeasurement.value_numeric.isnot(None),
            )
            .group_by(PalletMeasurement.parameter_name, PalletMeasurement.unit)
            .having(func.count(PalletMeasurement.id) >= 2)
            .order_by(func.count(PalletMeasurement.id).desc())
            .all()
        )

    # Grower performance
    grower_rows = []
    if load_ids:
        grower_rows = (
            db.query(
                Pallet.grower_code,
                func.count(Pallet.id).label("total"),
                func.sum(case((Pallet.status == "PASS", 1), else_=0)).label("passed"),
                func.sum(case((Pallet.status.in_(["FAIL", "REJECT"]), 1), else_=0)).label("failed"),
                func.sum(case((Pallet.status == "HOLD", 1), else_=0)).label("held"),
            )
            .filter(
                Pallet.load_id.in_(load_ids),
                Pallet.grower_code.isnot(None), Pallet.grower_code != "",
                Pallet.status.in_(["PASS", "FAIL", "REJECT", "HOLD"]),
            )
            .group_by(Pallet.grower_code)
            .order_by(func.count(Pallet.id).desc())
            .limit(15)
            .all()
        )

    # ── Build PDF ──────────────────────────────────────────────────────────────
    out_dir = Path(settings.EXPORT_DIR) / "analytics"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"analytics_{uuid.uuid4().hex[:8]}.pdf"

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    brand  = colors.Color(*[c / 255 for c in _PURPLE])
    pass_c = colors.Color(*[c / 255 for c in _PASS])
    hold_c = colors.Color(*[c / 255 for c in _HOLD])
    rej_c  = colors.Color(*[c / 255 for c in _REJECT])

    title_style = ParagraphStyle("title", parent=styles["Heading1"],
                                 textColor=brand, fontSize=18, spaceAfter=4)
    h2_style    = ParagraphStyle("h2", parent=styles["Heading2"],
                                 textColor=brand, fontSize=11, spaceBefore=12, spaceAfter=4)
    small       = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, textColor=colors.grey)

    def _tbl(rows, col_widths=None, extra_styles=None):
        base = [
            ("BACKGROUND", (0, 0), (-1, 0), brand),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("GRID",       (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.98, 0.97, 0.99)]),
            ("ALIGN",      (1, 0), (-1, -1), "CENTER"),
            ("TOPPADDING",    (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        if extra_styles:
            base.extend(extra_styles)
        t = Table(rows, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle(base))
        return t

    def _rate_color(rate):
        if rate is None: return colors.grey
        if rate >= 85: return pass_c
        if rate >= 60: return hold_c
        return rej_c

    story = []

    # Title
    period_label = f"Last {months} months" if months < 36 else "All time"
    story.append(Paragraph("MAGOPCO", title_style))
    story.append(Paragraph("Quality Analytics Report", styles["Heading2"]))
    story.append(HRFlowable(width="100%", color=brand, thickness=2, spaceAfter=6))
    story.append(Paragraph(
        f"Period: {period_label} · Client: {client_name} · "
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC",
        small,
    ))
    story.append(Spacer(1, 0.4*cm))

    # KPI summary table
    story.append(Paragraph("Key Performance Indicators", h2_style))
    kpi_data = [
        ["Overall Pass Rate", "Avg Quality Score", "Loads Analysed", "Total Pallets", "Biggest Risk"],
        [
            f"{pass_rate}%" if pass_rate is not None else "—",
            str(round(float(avg_qs), 1)) if avg_qs else "—",
            str(len(load_ids)),
            str(total_pallets),
            f"{worst_param}\n({worst_rate}%)" if worst_param else "None",
        ],
    ]
    kpi_t = Table(kpi_data, colWidths=[3.6*cm]*5)
    kpi_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
        ("GRID",       (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("FONTNAME",   (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 1), (-1, 1), 12),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])
    if pass_rate is not None:
        kpi_style.add("TEXTCOLOR", (0, 1), (0, 1), _rate_color(pass_rate))
    if worst_param and worst_rate is not None and worst_rate < 70:
        kpi_style.add("TEXTCOLOR", (4, 1), (4, 1), rej_c)
    kpi_t.setStyle(kpi_style)
    story.append(kpi_t)
    story.append(Spacer(1, 0.3*cm))

    # Quality trends table
    if trend_rows:
        story.append(Paragraph("Quality Trends by Period", h2_style))
        t_rows = [["Period", "Pallets", "Passed", "Hold", "Failed", "Pass Rate %"]]
        extra = []
        for r in trend_rows:
            passed  = int(r.passed or 0)
            failed  = int(r.failed or 0)
            held    = int(r.held   or 0)
            total   = int(r.total  or 0)
            analysed = passed + failed + held
            pr = round(passed / analysed * 100, 1) if analysed else None
            row_i = len(t_rows)
            t_rows.append([
                str(r.period)[:10],
                str(total), str(passed), str(held), str(failed),
                f"{pr}%" if pr is not None else "—",
            ])
            if pr is not None:
                c_ = pass_c if pr >= 85 else (hold_c if pr >= 60 else rej_c)
                extra += [("TEXTCOLOR", (-1, row_i), (-1, row_i), c_),
                          ("FONTNAME",  (-1, row_i), (-1, row_i), "Helvetica-Bold")]
        story.append(_tbl(t_rows, col_widths=[3.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm], extra_styles=extra))
        story.append(Spacer(1, 0.3*cm))

    # Parameter compliance table
    if compliance_rows:
        story.append(Paragraph("Parameter Compliance", h2_style))
        c_rows = [["Parameter", "Unit", "Total", "Passed", "Compliance %", "Avg Value", "Std Max"]]
        extra = []
        for r in compliance_rows:
            total_r = int(r.total)
            passed_r = int(r.passed or 0)
            cr = round(passed_r / total_r * 100, 1) if total_r else None
            row_i = len(c_rows)
            c_rows.append([
                r.parameter_name,
                r.unit or "—",
                str(total_r),
                str(passed_r),
                f"{cr}%" if cr is not None else "—",
                str(round(float(r.avg_value), 2)) if r.avg_value else "—",
                str(round(float(r.std_max), 2)) if r.std_max else "—",
            ])
            if cr is not None:
                c_ = pass_c if cr >= 90 else (hold_c if cr >= 70 else rej_c)
                extra += [("TEXTCOLOR", (4, row_i), (4, row_i), c_),
                          ("FONTNAME",  (4, row_i), (4, row_i), "Helvetica-Bold")]
        story.append(_tbl(c_rows, col_widths=[4.5*cm, 1.5*cm, 1.5*cm, 1.8*cm, 2.5*cm, 2.2*cm, 2.2*cm], extra_styles=extra))
        story.append(Spacer(1, 0.3*cm))

    # Grower performance table
    if grower_rows:
        story.append(Paragraph("Grower Performance", h2_style))
        g_rows = [["Grower Code", "Total Pallets", "Passed", "Hold", "Failed", "Pass Rate %", "Fail Rate %"]]
        extra = []
        for r in grower_rows:
            p_ = int(r.passed or 0)
            f_ = int(r.failed or 0)
            h_ = int(r.held   or 0)
            t_ = int(r.total  or 0)
            an_ = p_ + f_ + h_
            pr_ = round(p_ / an_ * 100, 1) if an_ else None
            fr_ = round(f_ / an_ * 100, 1) if an_ else None
            row_i = len(g_rows)
            g_rows.append([
                r.grower_code or "—",
                str(t_), str(p_), str(h_), str(f_),
                f"{pr_}%" if pr_ is not None else "—",
                f"{fr_}%" if fr_ is not None else "—",
            ])
            if pr_ is not None:
                c_ = pass_c if pr_ >= 85 else (hold_c if pr_ >= 60 else rej_c)
                extra += [("TEXTCOLOR", (5, row_i), (5, row_i), c_),
                          ("FONTNAME",  (5, row_i), (5, row_i), "Helvetica-Bold")]
        story.append(_tbl(g_rows, col_widths=[3.5*cm, 2.8*cm, 2*cm, 2*cm, 2*cm, 2.3*cm, 2.3*cm], extra_styles=extra))

    # Footer
    story.append(Spacer(1, 0.5*cm))
    story.append(HRFlowable(width="100%", color=colors.lightgrey, spaceAfter=4))
    story.append(Paragraph(
        f"Magopco Quality Intelligence Platform · Analytics Report · "
        f"{datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC",
        small,
    ))

    doc.build(story)
    return out_path


# ── Shared loader ─────────────────────────────────────────────────────────────

def _load_with_relations(db: Session, load_id: int) -> Load:
    from fastapi import HTTPException
    load = (
        db.query(Load)
        .options(
            joinedload(Load.pallets).joinedload(Pallet.measurements),
            joinedload(Load.summary_measurements),
        )
        .filter(Load.id == load_id)
        .first()
    )
    if not load:
        raise HTTPException(status_code=404, detail="Load not found")
    return load
